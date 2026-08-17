from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

template = r"analyste-service/src/main/resources/templates/Template_Rapport_Audit_ProjectIQ.docx"
doc = Document(template)
template_table_style = doc.tables[0].style if doc.tables else None

# Remove the old plain-text finalisation section, if it was generated before.
start = next((i for i, p in enumerate(doc.paragraphs) if "Finalisation du Pack" in p.text), None)
if start is not None:
    # The two tables at the end belong to the previous finalisation section.
    for table in list(doc.tables[-2:]):
        table._element.getparent().remove(table._element)
    for paragraph in list(doc.paragraphs[start:]):
        paragraph._element.getparent().remove(paragraph._element)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def cell_text(cell, text, bold=False, color=None):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(14)
heading.paragraph_format.space_after = Pt(7)
run = heading.add_run("5. Finalisation du Pack et Decision Finale")
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(31, 78, 121)

summary = doc.add_table(rows=0, cols=2)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
if template_table_style:
    summary.style = template_table_style
for label, placeholder in [
    ("Documents du pack", "[[PACK_DOCUMENTS]]"),
    ("Envoi aux decideurs", "[[VALIDATION_SENT_AT]]"),
    ("Decision finale", "[[DECISION_FINALE]]"),
    ("Source de decision", "[[FINAL_DECISION_SOURCE]]"),
    ("Date et heure de decision", "[[FINAL_DECISION_AT]]"),
    ("Statut de cloture", "[[FINAL_STATUS]]"),
    ("Date et heure d'archivage", "[[ARCHIVED_AT]]"),
]:
    cells = summary.add_row().cells
    shade(cells[0], 'D9EAF7')
    cell_text(cells[0], label, bold=True, color=(31, 78, 121))
    cell_text(cells[1], placeholder)

subheading = doc.add_paragraph()
subheading.paragraph_format.space_before = Pt(12)
subheading.paragraph_format.space_after = Pt(6)
subrun = subheading.add_run("5.1 Validations des Decideurs")
subrun.bold = True
subrun.font.size = Pt(12)
subrun.font.color.rgb = RGBColor(31, 78, 121)

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
if template_table_style:
    table.style = template_table_style
headers = ["Manager", "Role", "Decision", "Source", "Date et heure", "Commentaire"]
for cell, text in zip(table.rows[0].cells, headers):
    shade(cell, 'D9EAF7')
    cell_text(cell, text, bold=True, color=(31, 78, 121))
row = table.add_row()
cell_text(row.cells[0], '[[VALIDATION_ROWS]]')
for cell in row.cells[1:]:
    cell_text(cell, '')

doc.save(template)
